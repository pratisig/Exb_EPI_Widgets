#!/usr/bin/env python3
"""Convert a repository Markdown document into printable HTML, PDF and DOCX.

Used to publish the GIS Centre documents (proposal dossier, user manual) as files
that can be attached to a ticket or printed for a meeting.

Requirements: Python 3.9+, `markdown`, `xhtml2pdf`, `python-docx`.
    python -m pip install markdown xhtml2pdf python-docx

Usage:
    python tools/md-to-deliverables.py DOSSIER_PROPOSITION_GIS_CENTRE.md deliverables
    python tools/md-to-deliverables.py README.md deliverables DESCRIPTION_OUTIL_EPI_AGGREGATOR

The optional third argument renames the generated files, which is useful when the
Markdown file name is not the name to publish (e.g. README.md as "Description de l'outil").
"""
import html
import os
import re
import sys

try:
    import markdown
    from xhtml2pdf import pisa
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm, RGBColor
except ImportError as exc:  # pragma: no cover - dependency hint
    sys.exit(f"Dépendance manquante ({exc}). Exécuter : python -m pip install markdown xhtml2pdf python-docx")

PRINT_CSS = """
@page { size: A4; margin: 1.6cm 1.4cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 9.5pt; color: #1f2933; line-height: 1.45; }
h1 { font-size: 18pt; color: #0b4f83; border-bottom: 2px solid #0b4f83; padding-bottom: 6px; }
h2 { font-size: 13pt; color: #1261a0; margin-top: 16px; border-bottom: 1px solid #cbd5e1; padding-bottom: 3px; }
h3 { font-size: 11pt; color: #1d4f6e; margin-top: 12px; }
h4 { font-size: 10pt; color: #1d4f6e; }
table { border-collapse: collapse; width: 100%; margin: 6px 0 10px 0; font-size: 8pt; }
th, td { border: 1px solid #b8c4d0; padding: 4px 5px; text-align: left; vertical-align: top; word-wrap: break-word; }
th { background: #eaf4fa; color: #0b4f83; }
code, pre { font-family: Courier, monospace; font-size: 8pt; background: #f4f6f8; }
pre { border: 1px solid #dde3e9; padding: 6px; white-space: pre-wrap; }
blockquote { border-left: 3px solid #1261a0; margin: 6px 0; padding: 2px 8px; background: #f5fafd; }
ul, ol { margin: 4px 0 8px 18px; }
li { margin-bottom: 2px; }
hr { border: 0; border-top: 1px solid #cbd5e1; margin: 14px 0; }
"""


def md_to_html(text: str, title: str) -> str:
    body = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists", "attr_list", "toc"],
        extension_configs={"toc": {"title": ""}},
    )
    return (
        "<!DOCTYPE html><html lang='fr'><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title><style>{PRINT_CSS}</style></head>"
        f"<body>{body}</body></html>"
    )


def html_to_pdf(html_doc: str, pdf_path: str) -> None:
    with open(pdf_path, "wb") as handle:
        result = pisa.CreatePDF(html_doc, dest=handle, encoding="utf-8")
    if result.err:
        raise RuntimeError(f"Échec de génération du PDF ({result.err} erreur(s))")


# --------------------------------------------------------------------------- DOCX

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text: str) -> None:
    """Render **bold** and `code` inline spans as Word runs."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        else:
            paragraph.add_run(part)


def split_row(line: str):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def md_to_docx(text: str, docx_path: str, title: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    lines = text.split("\n")
    index = 0
    first_heading_skipped = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = doc.add_paragraph()
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue

        # table
        if stripped.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].strip()) <= set("|-: "):
            headers = split_row(stripped)
            rows = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            for position, header in enumerate(headers):
                cell = table.rows[0].cells[position]
                cell.text = ""
                add_runs(cell.paragraphs[0], header)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for position, value in enumerate(row[: len(headers)]):
                    cells[position].text = ""
                    add_runs(cells[position].paragraphs[0], value)
            continue

        # headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            content = stripped[level:].strip()
            if level == 1 and not first_heading_skipped:
                first_heading_skipped = True
                heading = doc.add_heading(content, level=0)
                heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                doc.add_heading(content, level=min(level, 4))
            index += 1
            continue

        # horizontal rule
        if stripped in ("---", "***", "___"):
            doc.add_paragraph("_" * 60)
            index += 1
            continue

        # bullets
        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:])
            index += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            add_runs(paragraph, stripped.lstrip("> ").strip())
            paragraph.paragraph_format.left_indent = Cm(0.6)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x1D, 0x4F, 0x6E)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_runs(paragraph, stripped)
        index += 1

    doc.save(docx_path)


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    source = os.path.abspath(sys.argv[1])
    outdir = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.join(os.path.dirname(source), "deliverables")
    os.makedirs(outdir, exist_ok=True)

    with open(source, encoding="utf-8") as handle:
        text = handle.read()
    title = next((line.lstrip("# ").strip() for line in text.split("\n") if line.startswith("# ")), os.path.basename(source))
    base = sys.argv[3].strip() if len(sys.argv) > 3 and sys.argv[3].strip() else os.path.splitext(os.path.basename(source))[0]

    html_doc = md_to_html(text, title)
    html_path = os.path.join(outdir, f"{base}.html")
    pdf_path = os.path.join(outdir, f"{base}.pdf")
    docx_path = os.path.join(outdir, f"{base}.docx")
    with open(html_path, "w", encoding="utf-8") as handle:
        handle.write(html_doc)
    html_to_pdf(html_doc, pdf_path)
    md_to_docx(text, docx_path, title)

    for path in (html_path, pdf_path, docx_path):
        print(f"{os.path.relpath(path)}  ({os.path.getsize(path) // 1024} Ko)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
